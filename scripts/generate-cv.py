#!/usr/bin/env python3
"""Generate Renet Ruseva CV as Word document from portfolio content."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = "/Users/renetruseva/Documents/portfolio2026/renet-ruseva-cv.docx"


def set_narrow_margins(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1E, 0x50, 0xE1)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    # bottom border
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1E50E1")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.15)
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text)
    else:
        p.add_run(text)


def add_project(doc, title, subtitle, bullets, tech):
    p = doc.add_paragraph()
    title_run = p.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(10.5)
    if subtitle:
        sub = p.add_run(f"  —  {subtitle}")
        sub.italic = True
        sub.font.size = Pt(10)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    for b in bullets:
        add_bullet(doc, b)
    tech_p = doc.add_paragraph()
    tech_run = tech_p.add_run(f"Tools: {tech}")
    tech_run.italic = True
    tech_run.font.size = Pt(9)
    tech_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    tech_p.paragraph_format.space_after = Pt(4)


def main():
    doc = Document()
    set_narrow_margins(doc)

    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    # Header
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = name.add_run("RENET RUSEVA")
    nr.bold = True
    nr.font.size = Pt(22)
    nr.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("Creative Developer")
    tr.font.size = Pt(12)
    tr.font.color.rgb = RGBColor(0x1E, 0x50, 0xE1)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(6)
    parts = [
        "Kortrijk, Belgium",
        "renetruseva@gmail.com",
        "linkedin.com/in/renet-ruseva-3089682b9",
        "github.com/renetruseva-prog",
        "behance.net/renetruseva",
    ]
    cr = contact.add_run("  ·  ".join(parts))
    cr.font.size = Pt(9)
    cr.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # Summary
    add_section_heading(doc, "Profile")
    summary = doc.add_paragraph(
        "Digital design student and creative developer based in Kortrijk, Belgium. "
        "I combine visual storytelling with web technology to build experiences that feel "
        "inspiring, authentic, and intuitive. I work at the intersection of design, "
        "technology, and user interaction — taking complex ideas and turning them into "
        "clear, engaging digital products."
    )
    summary.paragraph_format.space_after = Pt(4)

    # Education
    add_section_heading(doc, "Education")
    edu_entries = [
        (
            "Sep 2024 – Present",
            "Devine – KASK Howest",
            "Digital Design and Development",
            "Kortrijk, Belgium",
        ),
        (
            "Sep 2018 – May 2023",
            'High School of Mathematics "Dr Petar Beron"',
            None,
            "Varna, Bulgaria",
        ),
    ]
    for period, institution, program, location in edu_entries:
        p = doc.add_paragraph()
        r1 = p.add_run(f"{period}  ")
        r1.bold = True
        r1.font.size = Pt(10)
        r2 = p.add_run(institution)
        r2.bold = True
        if program:
            p.add_run(f" — {program}")
        loc = doc.add_paragraph()
        lr = loc.add_run(location)
        lr.italic = True
        lr.font.size = Pt(9.5)
        loc.paragraph_format.space_after = Pt(4)

    # Skills
    add_section_heading(doc, "Skills")
    skills = [
        ("Visual Design & UX/UI:", "Figma, Photoshop, Illustrator, After Effects, Miro, Prototyping"),
        ("Web Development:", "HTML, CSS, JavaScript, React, React Router, GSAP, Git, MySQL, ml5.js"),
        ("Additional:", "WebRTC, Socket.IO, Node.js, Express, Supabase, Leaflet, CI/CD, Mobile-first development"),
    ]
    for label, items in skills:
        p = doc.add_paragraph()
        lr = p.add_run(label + " ")
        lr.bold = True
        p.add_run(items)
        p.paragraph_format.space_after = Pt(3)

    # Projects
    add_section_heading(doc, "Selected Projects")
    projects = [
        (
            "MemoMe – Discover Antwerp",
            "Team project · Lead Developer, UX/UI Researcher",
            [
                "Built a location-based web app encouraging young travellers (18–36) to explore Antwerp through map-based memory pinning, QR sticker collectibles, and shareable travel diaries.",
                "Led development of core features: memo CRUD with media upload, geocoding, authentication (guest + registered), travel journals, and GSAP animations.",
                "Implemented React Router v7 architecture with Supabase backend and Leaflet mapping integration; followed CI/CD pipeline practices.",
            ],
            "React, React Router v7, JavaScript, CSS, GSAP, Supabase, Leaflet, Figma, Miro",
        ),
        (
            "The Beauty of Transgression",
            "Solo project · MoMu Museum, Antwerp",
            [
                "Designed and developed an interactive online experience immersing readers in a sensory journey about Belgian designer Walter van Beirendonck.",
                "Focused on interactivity, micro-animations, GSAP integration, and immersive storytelling.",
            ],
            "Figma, Photoshop, HTML, CSS, JavaScript, GSAP",
        ),
        (
            "WebRTC Presenter",
            "Solo project",
            [
                "Built a real-time presentation system turning a phone into a private presenter view with speaker notes, timer, slide navigation, laser pointer, and live camera.",
                "Implemented P2P WebRTC connection with Socket.IO signaling, QR code pairing over local network, and Reveal.js slide deck integration.",
            ],
            "JavaScript, WebRTC, Node.js, Express, Socket.IO, Reveal.js, HTML, CSS",
        ),
        (
            "Miles and Meals",
            "Solo project · Branding case",
            [
                "Created a full branding identity and app interface for a subscription food box service, including tone of voice, logo, colour system, typography, packaging, and dedicated app design.",
            ],
            "Figma, Photoshop, Adobe Illustrator",
        ),
    ]
    for title, subtitle, bullets, tech in projects:
        add_project(doc, title, subtitle, bullets, tech)

    # Languages
    add_section_heading(doc, "Languages")
    lang_p = doc.add_paragraph()
    lang_p.add_run("Bulgarian (Native)  ·  English (Fluent)  ·  Russian (Fluent)")

    # Interests
    add_section_heading(doc, "Interests")
    int_p = doc.add_paragraph(
        "Crocheting, reading, collaging, DIY, travelling, exploring cafés"
    )

    doc.save(OUTPUT)
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
